# -*- coding: utf-8 -*-
"""Generate docs4 by splitting docs2 into V1/V2 planning scopes.

V1 source:
    lastspr/v1checklist.xlsx
    sheet "3.CRM (admin)_제로스트_0518논의"
    rows 4..751

V2 source:
    rows 753..end plus docs2 sections that are not mapped to V1.

This script is intentionally conservative. It does not rewrite planning content;
it moves whole markdown sections into V1 or V2 and records the matching basis.
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DOCS2 = ROOT / "docs2"
DOCS4 = ROOT / "docs4"
EXCEL = ROOT / "lastspr" / "v1checklist.xlsx"
SHEET_NAME = "3.CRM (admin)_제로스트_0518논의"

V1_START_ROW = 4
V1_END_ROW = 751
V2_START_ROW = 753

EXCLUDED_DOCS2_TOP_DIRS = {
    "_공통",
    "registry",
    "reports",
    "scripts",
    "정합성체크",
}

# Examples matched:
# CLS-01-07, SAL-EXT-05-06, PAY-STF-02-09, SCR-M001, DLG-S016, IoT-05-04.
CODE_RE = re.compile(r"\b[A-Z][A-Z0-9]*(?:-[A-Z0-9]+){1,6}\b", re.IGNORECASE)
SECTION_RE = re.compile(r"(?m)^##\s+(.+?)\s*$")


@dataclass(frozen=True)
class ChecklistRecord:
    row: int
    contract_status: str
    level: str
    category: str
    contract_code: str
    contract_name: str
    planning_code: str
    planning_name: str
    dev_rr: str
    zerost_rr: str
    opinion: str
    client_feedback: str
    extra_opinion: str
    codes: tuple[str, ...]


@dataclass(frozen=True)
class Section:
    source_path: Path
    rel_path: Path
    title: str
    start_line: int
    content: str
    heading_codes: tuple[str, ...]
    body_codes: tuple[str, ...]
    ordered_codes: tuple[str, ...]


@dataclass
class SectionDecision:
    section: Section
    output_scope: str
    classification: str
    matched_v1_codes: tuple[str, ...]
    matched_v2_codes: tuple[str, ...]
    basis: str
    warning: str
    output_path: Path


def normalize_code(value: str) -> str:
    return value.strip().upper()


def extract_codes(value: object) -> tuple[str, ...]:
    if value is None:
        return ()
    text = str(value)
    found: list[str] = []
    seen: set[str] = set()
    for match in CODE_RE.findall(text):
        code = normalize_code(match)
        if code not in seen:
            seen.add(code)
            found.append(code)
    return tuple(found)


def unique_ordered(values: Iterable[str]) -> tuple[str, ...]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        normalized = normalize_code(value)
        if normalized and normalized not in seen:
            seen.add(normalized)
            ordered.append(normalized)
    return tuple(ordered)


def code_related(a: str, b: str) -> bool:
    """Return True when codes are exact or parent/child related."""
    a = normalize_code(a)
    b = normalize_code(b)
    return a == b or a.startswith(f"{b}-") or b.startswith(f"{a}-")


def related_matches(code: str, scope_codes: set[str]) -> tuple[str, ...]:
    return tuple(sorted(candidate for candidate in scope_codes if code_related(code, candidate)))


def matches_any(code: str, scope_codes: set[str]) -> bool:
    return any(code_related(code, candidate) for candidate in scope_codes)


def safe_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def load_records() -> tuple[list[ChecklistRecord], list[ChecklistRecord]]:
    # Random cell access is extremely slow in openpyxl read_only mode because
    # each ws.cell lookup can rescan the stream. This workbook is small enough
    # to load normally, and normal mode keeps the generation deterministic.
    wb = load_workbook(EXCEL, data_only=True, read_only=False)
    ws = wb[SHEET_NAME]

    def make_record(row: int) -> ChecklistRecord:
        values = [safe_cell(ws.cell(row, col).value) for col in range(1, 13)]
        codes = unique_ordered((*extract_codes(values[3]), *extract_codes(values[5])))
        return ChecklistRecord(
            row=row,
            contract_status=values[0],
            level=values[1],
            category=values[2],
            contract_code=values[3],
            contract_name=values[4],
            planning_code=values[5],
            planning_name=values[6],
            dev_rr=values[7],
            zerost_rr=values[8],
            opinion=values[9],
            client_feedback=values[10],
            extra_opinion=values[11],
            codes=codes,
        )

    v1_records = [make_record(row) for row in range(V1_START_ROW, V1_END_ROW + 1)]
    v2_records = [make_record(row) for row in range(V2_START_ROW, ws.max_row + 1)]
    return v1_records, v2_records


def records_to_code_map(records: Iterable[ChecklistRecord]) -> dict[str, list[ChecklistRecord]]:
    code_map: dict[str, list[ChecklistRecord]] = defaultdict(list)
    for record in records:
        for code in record.codes:
            code_map[code].append(record)
    return dict(code_map)


def docs2_source_files() -> list[Path]:
    files: list[Path] = []
    for path in sorted(DOCS2.rglob("*.md")):
        rel = path.relative_to(DOCS2)
        first = rel.parts[0]
        if first in EXCLUDED_DOCS2_TOP_DIRS:
            continue
        if first.startswith("D") and len(rel.parts) >= 2:
            files.append(path)
    return files


def line_number_at(text: str, index: int) -> int:
    return text.count("\n", 0, index) + 1


def split_sections(path: Path) -> list[Section]:
    text = path.read_text(encoding="utf-8")
    rel = path.relative_to(DOCS2)
    matches = list(SECTION_RE.finditer(text))
    sections: list[Section] = []

    def build_section(title: str, start: int, end: int) -> Section:
        content = text[start:end].strip()
        heading_codes = extract_codes(title)
        body_codes = extract_codes(content)
        ordered_codes = unique_ordered((*extract_codes(title), *extract_codes(content)))
        return Section(
            source_path=path,
            rel_path=rel,
            title=title,
            start_line=line_number_at(text, start),
            content=content,
            heading_codes=heading_codes,
            body_codes=body_codes,
            ordered_codes=ordered_codes,
        )

    if not matches:
        sections.append(build_section("문서 전체", 0, len(text)))
        return sections

    preamble = text[: matches[0].start()].strip()
    if preamble:
        sections.append(build_section("문서 서문", 0, matches[0].start()))

    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append(build_section(match.group(1), start, end))
    return sections


def is_support_section(section: Section) -> bool:
    title = section.title
    support_keywords = (
        "문서 서문",
        "도메인 개요",
        "핵심 원칙",
        "공통",
        "권한",
        "상태",
        "운영 정책",
        "운영정책",
        "자동화",
        "알림",
        "정책",
        "KPI",
        "외부 연동",
        "예외",
        "토스트",
        "용어",
        "기준",
    )
    return any(keyword in title for keyword in support_keywords)


def first_scope_from_ordered_codes(
    ordered_codes: Iterable[str], v1_codes: set[str], v2_codes: set[str]
) -> str:
    for code in ordered_codes:
        if matches_any(code, v1_codes):
            return "V1"
        if matches_any(code, v2_codes):
            return "V2"
    return "V2"


def format_record(record: ChecklistRecord) -> str:
    name = record.planning_name or record.contract_name
    code = record.planning_code or record.contract_code
    if name:
        return f"{record.row}행 `{code}` {name}"
    return f"{record.row}행 `{code}`"


def format_basis(
    matched_codes: Iterable[str],
    code_map: dict[str, list[ChecklistRecord]],
    fallback: str,
) -> str:
    codes = tuple(matched_codes)
    if not codes:
        return fallback

    parts: list[str] = []
    for code in codes[:6]:
        related_records: list[ChecklistRecord] = []
        for candidate, records in code_map.items():
            if code_related(code, candidate):
                related_records.extend(records)
        if related_records:
            parts.append(f"`{code}`: {format_record(related_records[0])}")
        else:
            parts.append(f"`{code}`")
    if len(codes) > 6:
        parts.append(f"외 {len(codes) - 6}개")
    return "; ".join(parts)


def decide_sections(
    all_sections: list[Section],
    v1_codes: set[str],
    effective_v2_codes: set[str],
    v1_code_map: dict[str, list[ChecklistRecord]],
    v2_code_map: dict[str, list[ChecklistRecord]],
) -> list[SectionDecision]:
    grouped: dict[Path, list[Section]] = defaultdict(list)
    for section in all_sections:
        grouped[section.rel_path].append(section)

    decisions: list[SectionDecision] = []
    for rel_path, sections in grouped.items():
        raw: list[tuple[Section, tuple[str, ...], tuple[str, ...], str]] = []
        file_has_v1 = False
        file_has_v2 = False

        for section in sections:
            matched_v1 = unique_ordered(
                match for code in section.ordered_codes for match in related_matches(code, v1_codes)
            )
            matched_v2 = unique_ordered(
                match for code in section.ordered_codes for match in related_matches(code, effective_v2_codes)
            )
            if matched_v1:
                file_has_v1 = True
            if matched_v2:
                file_has_v2 = True
            primary = first_scope_from_ordered_codes(section.ordered_codes, v1_codes, effective_v2_codes)
            raw.append((section, matched_v1, matched_v2, primary))

        for section, matched_v1, matched_v2, primary in raw:
            output_scopes: list[tuple[str, str, tuple[str, ...], tuple[str, ...], str]] = []
            warning = ""

            if matched_v1 and matched_v2:
                warning = "V1/V2 코드가 같은 섹션에 함께 존재합니다. 1차 개발 범위 우선 원칙과 첫 매칭 코드 기준으로 배치했습니다."
                if primary == "V1":
                    output_scopes.append(("V1", "V1_MIXED", matched_v1, matched_v2, warning))
                else:
                    output_scopes.append(("V2", "V2_MIXED", matched_v1, matched_v2, warning))
            elif matched_v1:
                output_scopes.append(("V1", "V1", matched_v1, (), ""))
            elif matched_v2:
                output_scopes.append(("V2", "V2", (), matched_v2, ""))
            elif is_support_section(section) and file_has_v1:
                support_warning = "직접 매칭 코드는 없지만 같은 원본 파일의 V1 구현 이해에 필요한 공통/운영 보조 섹션입니다."
                output_scopes.append(("V1", "V1_SUPPORT", (), (), support_warning))
                if file_has_v2:
                    output_scopes.append(("V2", "V2_SUPPORT", (), (), "V2 구현에도 공통으로 필요한 보조 섹션입니다."))
            else:
                output_scopes.append(("V2", "V2_UNMAPPED", (), (), "체크리스트 V1 코드와 직접 매칭되지 않아 V2/보류 영역으로 배치했습니다."))

            for scope, classification, out_v1, out_v2, out_warning in output_scopes:
                fallback = (
                    "직접 매칭 코드 없음. 같은 원본 파일의 구현 기준을 이해하기 위한 보조 섹션입니다."
                    if "SUPPORT" in classification
                    else "직접 매칭 코드 없음. V1 행 범위에 포함되지 않은 기타/보류 섹션입니다."
                )
                if scope == "V1":
                    basis = format_basis(out_v1, v1_code_map, fallback)
                else:
                    basis = format_basis(out_v2, v2_code_map, fallback)
                output_path = DOCS4 / scope / rel_path
                decisions.append(
                    SectionDecision(
                        section=section,
                        output_scope=scope,
                        classification=classification,
                        matched_v1_codes=out_v1,
                        matched_v2_codes=out_v2,
                        basis=basis,
                        warning=out_warning,
                        output_path=output_path,
                    )
                )

    return decisions


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fp:
        writer = csv.DictWriter(fp, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def records_csv_rows(records: list[ChecklistRecord]) -> list[dict[str, object]]:
    return [
        {
            "row": record.row,
            "contract_status": record.contract_status,
            "level": record.level,
            "category": record.category,
            "contract_code": record.contract_code,
            "contract_name": record.contract_name,
            "planning_code": record.planning_code,
            "planning_name": record.planning_name,
            "dev_rr": record.dev_rr,
            "zerost_rr": record.zerost_rr,
            "opinion": record.opinion,
            "client_feedback": record.client_feedback,
            "extra_opinion": record.extra_opinion,
            "extracted_codes": ", ".join(record.codes),
        }
        for record in records
    ]


def section_note(decision: SectionDecision) -> str:
    lines = [
        "> **docs4 Scope**: " + decision.output_scope,
        "> **분류**: `" + decision.classification + "`",
        "> **분류 근거**: " + decision.basis,
    ]
    if decision.warning:
        lines.append("> **주의**: " + decision.warning)
    return "\n".join(lines) + "\n\n"


def write_scope_docs(decisions: list[SectionDecision]) -> None:
    by_output: dict[Path, list[SectionDecision]] = defaultdict(list)
    for decision in decisions:
        by_output[decision.output_path].append(decision)

    for output_path, output_decisions in sorted(by_output.items()):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        source_rel = output_decisions[0].section.rel_path.as_posix()
        scope = output_decisions[0].output_scope
        title = output_path.stem
        header = [
            f"# {title}",
            "",
            "> **docs4 Scope 문서**",
            f"> - 원본: `docs2/{source_rel}`",
            f"> - 범위: `{scope}`",
            f"> - 생성 기준: `{EXCEL.relative_to(ROOT).as_posix()}` `{SHEET_NAME}` 시트, V1=`{V1_START_ROW}~{V1_END_ROW}`행",
            "> - 원칙: 원본 기획 내용을 임의 축소하지 않고, 섹션 단위로 V1/V2에 재배치했습니다.",
            "",
        ]
        body_parts: list[str] = []
        for decision in output_decisions:
            body_parts.append(section_note(decision))
            body_parts.append(decision.section.content.strip())
            body_parts.append("\n")
        output_path.write_text("\n".join(header + body_parts).rstrip() + "\n", encoding="utf-8")


def write_indexes(decisions: list[SectionDecision]) -> None:
    for scope in ("V1", "V2"):
        rows = [decision for decision in decisions if decision.output_scope == scope]
        by_file: dict[Path, list[SectionDecision]] = defaultdict(list)
        for decision in rows:
            by_file[decision.output_path].append(decision)

        lines = [
            f"# docs4 {scope} INDEX",
            "",
            f"- 생성 기준: `{EXCEL.relative_to(ROOT).as_posix()}` `{SHEET_NAME}`",
            f"- V1 기준 행: `{V1_START_ROW}~{V1_END_ROW}`",
            f"- {scope} 출력 파일 수: {len(by_file)}",
            f"- {scope} 출력 섹션 수: {len(rows)}",
            "",
            "| 파일 | 섹션 수 | 분류 요약 |",
            "|---|---:|---|",
        ]
        for path, items in sorted(by_file.items()):
            rel = path.relative_to(DOCS4).as_posix()
            counts = Counter(item.classification for item in items)
            summary = ", ".join(f"{key} {value}" for key, value in sorted(counts.items()))
            lines.append(f"| `{rel}` | {len(items)} | {summary} |")
        (DOCS4 / f"{scope}_INDEX.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_readme(
    v1_records: list[ChecklistRecord],
    raw_v2_records: list[ChecklistRecord],
    decisions: list[SectionDecision],
    v1_codes: set[str],
    effective_v2_codes: set[str],
    v2_overrides: list[dict[str, object]],
) -> None:
    counts = Counter(decision.classification for decision in decisions)
    lines = [
        "# docs4",
        "",
        "docs4는 `docs2` 기획서를 개발 범위 기준으로 재배치한 납품용 문서입니다.",
        "",
        "## 생성 기준",
        "",
        f"- 체크리스트: `{EXCEL.relative_to(ROOT).as_posix()}`",
        f"- 시트: `{SHEET_NAME}`",
        f"- V1: {V1_START_ROW}행부터 {V1_END_ROW}행까지",
        f"- V2: {V2_START_ROW}행 이후 및 V1 코드와 직접 매칭되지 않는 보류/기타 섹션",
        "- V1/V2가 같은 섹션에 섞인 경우 V1 우선 원칙과 첫 매칭 코드 기준으로 배치하고 `_scope/section_scope_manifest.csv`에 표시했습니다.",
        "- `_공통`은 V1/V2 양쪽에서 참조하는 공통 정책/정의입니다.",
        "",
        "## 산출물",
        "",
        "- `V1/`: 1차 개발 범위 문서",
        "- `V2/`: 2차 이후 범위 또는 V1 직접 매칭이 없는 문서",
        "- `_공통/`: 공통 권한, 알림, KPI, NFR 등 공유 정의",
        "- `_scope/`: 체크리스트 원본 추출, 섹션 매핑, 검증 리포트",
        "",
        "## 요약",
        "",
        f"- V1 체크리스트 행 수: {len(v1_records)}",
        f"- V1 고유 코드 수: {len(v1_codes)}",
        f"- V2 원본 체크리스트 행 수: {len(raw_v2_records)}",
        f"- V2 유효 고유 코드 수: {len(effective_v2_codes)}",
        f"- V1 우선 원칙으로 V2에서 제외한 코드 수: {len(v2_overrides)}",
        f"- 섹션 분류 요약: {', '.join(f'{key} {value}' for key, value in sorted(counts.items()))}",
        "",
        "## 사용 방법",
        "",
        "1. V1 개발자는 `docs4/V1`과 `docs4/_공통`을 우선 봅니다.",
        "2. 섹션 상단의 `docs4 Scope` 메타 블록에서 엑셀 행과 매칭 근거를 확인합니다.",
        "3. 매칭 근거가 `V1_SUPPORT`, `V2_UNMAPPED`, `V1_MIXED`, `V2_MIXED`인 항목은 `_scope/scope_verification_report.md`에서 추가 확인합니다.",
        "",
    ]
    (DOCS4 / "README.md").write_text("\n".join(lines), encoding="utf-8")


def copy_common_docs() -> None:
    source_common = DOCS2 / "_공통"
    target_common = DOCS4 / "_공통"
    if source_common.exists():
        shutil.copytree(source_common, target_common)

    unresolved = DOCS2 / "_정책미확정_확인필요.md"
    if unresolved.exists():
        shutil.copy2(unresolved, DOCS4 / "_정책미확정_확인필요.md")


def write_scope_artifacts(
    v1_records: list[ChecklistRecord],
    raw_v2_records: list[ChecklistRecord],
    v1_codes: set[str],
    raw_v2_codes: set[str],
    effective_v2_codes: set[str],
    decisions: list[SectionDecision],
    all_sections: list[Section],
) -> None:
    scope_dir = DOCS4 / "_scope"
    scope_dir.mkdir(parents=True, exist_ok=True)

    record_fields = [
        "row",
        "contract_status",
        "level",
        "category",
        "contract_code",
        "contract_name",
        "planning_code",
        "planning_name",
        "dev_rr",
        "zerost_rr",
        "opinion",
        "client_feedback",
        "extra_opinion",
        "extracted_codes",
    ]
    write_csv(scope_dir / "v1_checklist_rows_4_751.csv", records_csv_rows(v1_records), record_fields)
    write_csv(scope_dir / "v2_explicit_rows_753_end.csv", records_csv_rows(raw_v2_records), record_fields)

    v2_overrides: list[dict[str, object]] = []
    for code in sorted(raw_v2_codes - effective_v2_codes):
        related_v1 = sorted(candidate for candidate in v1_codes if code_related(code, candidate))
        v2_overrides.append(
            {
                "v2_code": code,
                "related_v1_codes": ", ".join(related_v1),
                "reason": "V1 행 범위에 동일/상위/하위 코드가 있어 V1 우선 원칙으로 V2 매칭에서 제외",
            }
        )
    write_csv(
        scope_dir / "v2_rows_overridden_by_v1_priority.csv",
        v2_overrides,
        ["v2_code", "related_v1_codes", "reason"],
    )

    manifest_rows: list[dict[str, object]] = []
    for decision in decisions:
        section = decision.section
        manifest_rows.append(
            {
                "source_path": f"docs2/{section.rel_path.as_posix()}",
                "output_scope": decision.output_scope,
                "classification": decision.classification,
                "title": section.title,
                "start_line": section.start_line,
                "heading_codes": ", ".join(section.heading_codes),
                "body_codes": ", ".join(section.body_codes),
                "matched_v1_codes": ", ".join(decision.matched_v1_codes),
                "matched_v2_codes": ", ".join(decision.matched_v2_codes),
                "output_path": decision.output_path.relative_to(DOCS4).as_posix(),
                "basis": decision.basis,
                "warning": decision.warning,
            }
        )
    write_csv(
        scope_dir / "section_scope_manifest.csv",
        manifest_rows,
        [
            "source_path",
            "output_scope",
            "classification",
            "title",
            "start_line",
            "heading_codes",
            "body_codes",
            "matched_v1_codes",
            "matched_v2_codes",
            "output_path",
            "basis",
            "warning",
        ],
    )

    section_codes = set(code for section in all_sections for code in section.ordered_codes)
    unmatched_v1 = [
        code for code in sorted(v1_codes) if not any(code_related(code, section_code) for section_code in section_codes)
    ]
    unmatched_v2 = [
        code
        for code in sorted(effective_v2_codes)
        if not any(code_related(code, section_code) for section_code in section_codes)
    ]
    mixed = [decision for decision in decisions if decision.classification in {"V1_MIXED", "V2_MIXED"}]
    unmapped = [decision for decision in decisions if decision.classification == "V2_UNMAPPED"]
    support = [decision for decision in decisions if "SUPPORT" in decision.classification]

    counts = Counter(decision.classification for decision in decisions)
    lines = [
        "# docs4 Scope Verification Report",
        "",
        "## 기준",
        "",
        f"- V1 체크리스트: `{EXCEL.relative_to(ROOT).as_posix()}` `{SHEET_NAME}` {V1_START_ROW}~{V1_END_ROW}행",
        f"- V2 명시 항목: 같은 시트 {V2_START_ROW}행 이후",
        "- V1 우선 원칙: 같은 코드 또는 상위/하위 코드가 V1과 V2 양쪽에 있으면 V1로 우선 배치",
        "- 직접 매칭이 없는 일반 섹션은 V2_UNMAPPED로 보류하되, V1 구현 이해에 필요한 공통/운영 섹션은 V1_SUPPORT로 복사",
        "",
        "## 집계",
        "",
        f"- V1 고유 코드: {len(v1_codes)}",
        f"- V2 원본 고유 코드: {len(raw_v2_codes)}",
        f"- V2 유효 고유 코드: {len(effective_v2_codes)}",
        f"- V1 우선 제외 V2 코드: {len(v2_overrides)}",
        f"- 소스 섹션 수: {len(all_sections)}",
        f"- 출력 섹션 수: {len(decisions)}",
        f"- 분류 요약: {', '.join(f'{key} {value}' for key, value in sorted(counts.items()))}",
        f"- V1/V2 혼합 섹션: {len(mixed)}",
        f"- V2_UNMAPPED 섹션: {len(unmapped)}",
        f"- SUPPORT 복사 섹션: {len(support)}",
        "",
        "## 반드시 확인해야 하는 항목",
        "",
        "### 1. V1 체크리스트에는 있으나 docs2 섹션 코드로 직접 발견되지 않은 코드",
        "",
    ]
    if unmatched_v1:
        for code in unmatched_v1:
            lines.append(f"- `{code}`")
    else:
        lines.append("- 없음")

    lines.extend(["", "### 2. V1/V2가 같은 섹션에 섞인 항목", ""])
    if mixed:
        for decision in mixed[:200]:
            section = decision.section
            lines.append(
                f"- `{decision.classification}` `{section.rel_path.as_posix()}` L{section.start_line} `{section.title}` "
                f"/ V1=`{', '.join(decision.matched_v1_codes)}` V2=`{', '.join(decision.matched_v2_codes)}`"
            )
        if len(mixed) > 200:
            lines.append(f"- 외 {len(mixed) - 200}개")
    else:
        lines.append("- 없음")

    lines.extend(["", "### 3. V2_UNMAPPED 섹션", ""])
    if unmapped:
        for decision in unmapped[:200]:
            section = decision.section
            lines.append(f"- `{section.rel_path.as_posix()}` L{section.start_line} `{section.title}`")
        if len(unmapped) > 200:
            lines.append(f"- 외 {len(unmapped) - 200}개")
    else:
        lines.append("- 없음")

    lines.extend(["", "### 4. V2 유효 코드이나 docs2 섹션 코드로 직접 발견되지 않은 코드", ""])
    if unmatched_v2:
        for code in unmatched_v2:
            lines.append(f"- `{code}`")
    else:
        lines.append("- 없음")

    (scope_dir / "scope_verification_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_output() -> None:
    required = [
        DOCS4 / "README.md",
        DOCS4 / "V1_INDEX.md",
        DOCS4 / "V2_INDEX.md",
        DOCS4 / "_공통",
        DOCS4 / "_scope" / "section_scope_manifest.csv",
        DOCS4 / "_scope" / "scope_verification_report.md",
    ]
    missing = [path for path in required if not path.exists()]
    if missing:
        raise RuntimeError("필수 산출물 누락: " + ", ".join(str(path.relative_to(ROOT)) for path in missing))

    source_common_files = [path for path in (DOCS2 / "_공통").rglob("*") if path.is_file()]
    target_common_files = [path for path in (DOCS4 / "_공통").rglob("*") if path.is_file()]
    if len(source_common_files) != len(target_common_files):
        raise RuntimeError(f"_공통 파일 수 불일치: docs2={len(source_common_files)}, docs4={len(target_common_files)}")

    manifest_path = DOCS4 / "_scope" / "section_scope_manifest.csv"
    with manifest_path.open("r", encoding="utf-8-sig", newline="") as fp:
        rows = list(csv.DictReader(fp))
    if not rows:
        raise RuntimeError("section_scope_manifest.csv가 비어 있습니다.")

    missing_outputs = []
    for row in rows:
        output = DOCS4 / row["output_path"]
        if not output.exists():
            missing_outputs.append(row["output_path"])
    if missing_outputs:
        raise RuntimeError("manifest 출력 파일 누락: " + ", ".join(missing_outputs[:20]))

    sample_v1 = next((DOCS4 / row["output_path"] for row in rows if row["output_scope"] == "V1"), None)
    sample_v2 = next((DOCS4 / row["output_path"] for row in rows if row["output_scope"] == "V2"), None)
    for sample in (sample_v1, sample_v2):
        if sample and "docs4 Scope 문서" not in sample.read_text(encoding="utf-8"):
            raise RuntimeError(f"Scope 메타 블록 누락: {sample.relative_to(ROOT)}")


def generate(force: bool) -> None:
    if not DOCS2.exists():
        raise FileNotFoundError(DOCS2)
    if not EXCEL.exists():
        raise FileNotFoundError(EXCEL)

    if DOCS4.exists():
        if not force:
            raise RuntimeError("docs4가 이미 있습니다. 재생성하려면 --force를 사용하세요.")
        resolved_docs4 = DOCS4.resolve()
        resolved_root = ROOT.resolve()
        if resolved_docs4 == resolved_root or resolved_root not in resolved_docs4.parents:
            raise RuntimeError(f"안전하지 않은 삭제 대상: {resolved_docs4}")
        shutil.rmtree(DOCS4)

    DOCS4.mkdir(parents=True)

    v1_records, raw_v2_records = load_records()
    v1_code_map = records_to_code_map(v1_records)
    raw_v2_code_map = records_to_code_map(raw_v2_records)
    v1_codes = set(v1_code_map)
    raw_v2_codes = set(raw_v2_code_map)
    effective_v2_codes = {
        code for code in raw_v2_codes if not any(code_related(code, v1_code) for v1_code in v1_codes)
    }
    effective_v2_code_map = {code: raw_v2_code_map[code] for code in effective_v2_codes}

    files = docs2_source_files()
    all_sections: list[Section] = []
    for file in files:
        all_sections.extend(split_sections(file))

    decisions = decide_sections(
        all_sections=all_sections,
        v1_codes=v1_codes,
        effective_v2_codes=effective_v2_codes,
        v1_code_map=v1_code_map,
        v2_code_map=effective_v2_code_map,
    )

    copy_common_docs()
    write_scope_docs(decisions)

    v2_overrides = [
        {
            "v2_code": code,
            "related_v1_codes": ", ".join(sorted(candidate for candidate in v1_codes if code_related(code, candidate))),
            "reason": "V1 행 범위에 동일/상위/하위 코드가 있어 V1 우선 원칙으로 V2 매칭에서 제외",
        }
        for code in sorted(raw_v2_codes - effective_v2_codes)
    ]
    write_readme(v1_records, raw_v2_records, decisions, v1_codes, effective_v2_codes, v2_overrides)
    write_indexes(decisions)
    write_scope_artifacts(
        v1_records=v1_records,
        raw_v2_records=raw_v2_records,
        v1_codes=v1_codes,
        raw_v2_codes=raw_v2_codes,
        effective_v2_codes=effective_v2_codes,
        decisions=decisions,
        all_sections=all_sections,
    )
    validate_output()

    counts = Counter(decision.classification for decision in decisions)
    print("docs4 generated")
    print(f"source files: {len(files)}")
    print(f"source sections: {len(all_sections)}")
    print(f"output sections: {len(decisions)}")
    print(f"v1 codes: {len(v1_codes)}")
    print(f"effective v2 codes: {len(effective_v2_codes)}")
    print("classifications: " + ", ".join(f"{key}={value}" for key, value in sorted(counts.items())))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true", help="Delete and recreate docs4 safely.")
    args = parser.parse_args()
    generate(force=args.force)


if __name__ == "__main__":
    main()
